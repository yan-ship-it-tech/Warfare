#!/usr/bin/env python3
"""Build the sourcing workbook from the JSON export.

Optional companion to scripts/export-equipment-list.mjs — that script is
dependency-free and produces the CSV/Markdown/JSON; this one turns the JSON
into a single formatted .xlsx so the sourcing work can happen in one file
(frozen headers, filters, colour-coded status, blank fill-in columns).

    node scripts/export-equipment-list.mjs      # first — writes exports/
    python3 scripts/export-equipment-xlsx.py    # then — writes the workbook

    python3 scripts/export-equipment-xlsx.py --dump-manifest
        Re-derives docs/3d-model-sourcing-manifest.json from the Pass 16
        .xlsx. The JSON twin is what the .mjs exporter reads, so the
        dependency on openpyxl stays confined to this script.

Requires: openpyxl (not otherwise used by this repo — pip install openpyxl).
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
EXPORTS = ROOT / "exports"
MANIFEST_XLSX = ROOT / "docs/3d-model-sourcing-manifest.xlsx"
MANIFEST_JSON = ROOT / "docs/3d-model-sourcing-manifest.json"

HEADER_FILL = PatternFill("solid", fgColor="1F3A4D")
HEADER_FONT = Font(color="FFFFFF", bold=True)
STATUS_FILLS = {
    "Sourced": PatternFill("solid", fgColor="D8EDD5"),
    "Weak / Verify": PatternFill("solid", fgColor="FCEBC8"),
    "No free source": PatternFill("solid", fgColor="F7D4D0"),
    "": PatternFill("solid", fgColor="E6E6E6"),  # never searched
}
FILL_IN = PatternFill("solid", fgColor="FFFDF0")  # columns for the sourcer

# Column widths by name; anything unlisted gets a default.
WIDTHS = {
    "name": 34,
    "representative_system": 46,
    "short_role": 70,
    "group": 20,
    "category": 26,
    "asset_id": 34,
    "current_icon_image": 46,
    "current_gallery_images": 46,
    "prior_model_source": 52,
    "prior_model_notes": 46,
    "reference_links": 46,
    "found_model_url": 30,
    "found_image_urls": 30,
    "sourcing_notes": 34,
    "wikipedia": 46,
    "manufacturer": 24,
    "icon_image_status": 30,
}
FILL_IN_COLS = {
    "found_model_url",
    "found_model_license",
    "found_model_poly_count",
    "found_image_urls",
    "found_image_license",
    "sourcing_notes",
}


def dump_manifest() -> None:
    """Pass 16's manifest .xlsx -> a greppable, diffable JSON twin."""
    wb = load_workbook(MANIFEST_XLSX, data_only=True)
    out = {
        "source_file": "docs/3d-model-sourcing-manifest.xlsx",
        "note": "Machine-readable twin of the Pass 16 sourcing manifest. "
                "Regenerate with `python3 scripts/export-equipment-xlsx.py --dump-manifest`.",
        "readme": [],
        "rows": [],
    }
    for line in wb["README"].iter_rows(values_only=True):
        cells = [str(c).strip() for c in line if c not in (None, "")]
        if cells:
            out["readme"].append(" — ".join(cells))
    for sheet in ("Ukraine", "Russia"):
        rows = list(wb[sheet].iter_rows(values_only=True))
        hdr = [str(h or "").strip() for h in rows[0]]
        for row in rows[1:]:
            if not any(row):
                continue
            rec = {h: ("" if v is None else str(v).strip()) for h, v in zip(hdr, row)}
            rec["sheet"] = sheet
            out["rows"].append(rec)
    MANIFEST_JSON.write_text(json.dumps(out, indent=2, ensure_ascii=False) + "\n")
    print(f"wrote {MANIFEST_JSON.relative_to(ROOT)} ({len(out['rows'])} rows)")


def write_table(ws, rows: list[dict], status_col: str | None = "prior_sourcing_status") -> None:
    if not rows:
        return
    cols = list(rows[0].keys())
    ws.append(cols)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row in rows:
        ws.append([row.get(c, "") for c in cols])

    for idx, col in enumerate(cols, start=1):
        letter = get_column_letter(idx)
        ws.column_dimensions[letter].width = WIDTHS.get(col, 18)
        if col in FILL_IN_COLS:
            for r in range(2, len(rows) + 2):
                ws.cell(row=r, column=idx).fill = FILL_IN

    if status_col and status_col in cols:
        idx = cols.index(status_col) + 1
        for r, row in enumerate(rows, start=2):
            fill = STATUS_FILLS.get(row.get(status_col, ""))
            if fill:
                ws.cell(row=r, column=idx).fill = fill

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(rows) + 1}"
    ws.row_dimensions[1].height = 30


def readme_sheet(ws, data: dict) -> None:
    counts = data["counts"]
    assets = data["assets"]
    never = [a for a in assets if a["kind"] == "platform" and not a["prior_sourcing_status"]]
    lines = [
        ("Equipment sourcing export — Multi-Domain Warfare Digital Twin", True),
        ("", False),
        ("Generated by scripts/export-equipment-list.mjs + scripts/export-equipment-xlsx.py "
         "from the live roster (data/assets/*.json, data/catalog/*.json). Regenerate rather "
         "than hand-edit — a re-run overwrites this file.", False),
        ("", False),
        ("Sheets", True),
        (f"Ukraine — {counts['ukraine']} assets fielded by side_a.", False),
        (f"Russia — {counts['russia']} assets fielded by side_b.", False),
        (f"Reference Systems — {counts['catalog_systems']} swap-target systems from data/catalog/*.json. "
         "Shown in the per-asset comparison dropdown, so they need imagery too; they have no icon today.", False),
        ("", False),
        ("How to use it", True),
        ("Search on 'representative_system', not 'name' — the display name is sometimes a role "
         "('Point of Injury') while the representative system names the real hardware.", False),
        ("The pale yellow columns (found_model_url onward) are yours to fill in. Everything left "
         "of them is derived and will be overwritten on the next export.", False),
        ("'prior_*' columns carry the Pass 16 manifest's findings, matched by name: green = a free "
         "model was found, amber = a candidate exists but licence/variant was never verified, red = "
         "searched, nothing free, grey = never searched at all.", False),
        ("kind = 'composite/site' marks the ~9 entries that are a concept or a place rather than a "
         "named system (a logistics point, a grid node). Source reference photography of the type; "
         "there is no single real-world system to model.", False),
        ("has_hero_3d_model = yes marks the 11 assets that already have authored geometry in "
         "src/three/models.ts — a sourced model there is a replacement, not a gap.", False),
        ("", False),
        ("Where the work is", True),
        (f"{len(never)} named systems were never searched at all (grey status).", False),
        (f"{len([a for a in assets if a['prior_sourcing_status'] == 'No free source'])} were searched "
         "with no free source found — paid licence or commission.", False),
        (f"{len([a for a in assets if a['prior_sourcing_status'] == 'Weak / Verify'])} have an "
         "unverified candidate.", False),
        (f"Imagery is close to greenfield: "
         f"{len([a for a in assets if a['icon_image_status'].startswith('local placeholder')])} of the "
         f"{counts['assets']} icon_image values are /icons/<slug>.png placeholder paths for files that do not "
         f"exist in the repo, only {len([a for a in assets if a['icon_image_status'] == 'external url'])} point "
         f"at a real (Wikimedia Commons) image, and only "
         f"{len([a for a in assets if a['gallery_image_count'] > 0])} assets have any gallery image.", False),
        ("A sourced photograph belongs in gallery_images, which the detail panel renders — not in icon_image. "
         "The map glyph is the authored SVG set in src/icons/registry.tsx and deliberately never reads "
         "icon_image (a photo is unreadable at 32-44 px).", False),
        ("", False),
        ("Licensing", True),
        ("Binary assets cannot be fetched from an agent session (egress proxy blocks them), so every "
         "model/image here needs a human to download it and read its licence.", False),
        ("Anything sourced must carry visible attribution where it is rendered — the same obligation "
         "data/osm/ already carries for OpenStreetMap (ODbL).", False),
        ("CC BY-NC-ND is a trap for anything that may be reused commercially; the Pass 16 manifest "
         "flags at least one (Shahed-136).", False),
        ("Imported geometry still has to meet docs/MODEL_STYLE_GUIDE.md — a 190k-tri Orlan-10 needs "
         "decimating before it goes near the map view.", False),
    ]
    ws.column_dimensions["A"].width = 130
    for text, is_heading in lines:
        ws.append([text])
        cell = ws.cell(row=ws.max_row, column=1)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if is_heading:
            cell.font = Font(bold=True, size=13 if ws.max_row == 1 else 11)


def main() -> int:
    if "--dump-manifest" in sys.argv:
        dump_manifest()
        return 0

    src = EXPORTS / "equipment-export.json"
    if not src.exists():
        print(f"missing {src.relative_to(ROOT)} — run `node scripts/export-equipment-list.mjs` first")
        return 1
    data = json.loads(src.read_text())

    wb = Workbook()
    readme_sheet(wb.active, data)
    wb.active.title = "README"
    write_table(wb.create_sheet("Ukraine"), [a for a in data["assets"] if a["side_code"] == "UA"])
    write_table(wb.create_sheet("Russia"), [a for a in data["assets"] if a["side_code"] == "RU"])
    write_table(wb.create_sheet("Reference Systems"), data["catalog_systems"], status_col=None)

    dest = EXPORTS / "equipment-sourcing.xlsx"
    wb.save(dest)
    print(f"wrote {dest.relative_to(ROOT)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
